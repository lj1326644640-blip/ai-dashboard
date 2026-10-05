# -*- coding: utf-8 -*-
"""验收测试：字段完整性 / 数据准确性 / 看板HTML质量 / 线上一致性。

用法: python tests/acceptance_check.py   （自动取 data/ 下最新一天的归档）
"""
import glob
import json
import os
import re

import openpyxl

DAY_DIR = sorted(glob.glob("data/*/"))[-1].rstrip("/\\")
DATE = os.path.basename(DAY_DIR)
print(f"验收对象: {DAY_DIR} ({DATE})")

print()
print("====== 测试1: 统一字段且无重复 ======")
REQUIRED = ["id", "platform", "platform_name", "title", "url", "author", "published",
            "engagement", "engagement_total", "engagement_summary", "fingerprint",
            "keyword_hits", "s_engagement", "s_velocity", "s_relevance", "s_recency",
            "rule_score", "hours_display", "final_score", "why", "llm_relevance",
            "llm_value", "ctype"]
cand = json.loads(open(f"{DAY_DIR}/full.json", encoding="utf-8").read())
from collections import Counter  # noqa: E402
dup = {k: v for k, v in Counter(r["id"] for r in cand).items() if v > 1}
dup_url = {k: v for k, v in Counter(r["url"] for r in cand if r["url"]).items() if v > 1}
dup_fp = {k: v for k, v in Counter(r["fingerprint"] for r in cand).items() if v > 1}
missing = {r["id"]: [k for k in REQUIRED if k not in r or r[k] is None] for r in cand}
missing = {k: v for k, v in missing.items() if v}
print(f"候选 {len(cand)} 条 × {len(REQUIRED)} 个统一字段 | 缺失: {len(missing)}",
      "✅" if not missing else list(missing.items())[:3])
print(f"重复ID: {len(dup)} | 重复URL: {len(dup_url)} | 重复指纹: {len(dup_fp)}",
      "✅" if not (dup or dup_url or dup_fp) else "❌")

print()
print("====== 测试2: Top10 主题相关(关键词层) ======")
meta = json.loads(open(f"{DAY_DIR}/meta.json", encoding="utf-8").read())
weak = [i["rank"] for i in meta["items"][:10] if not i.get("ctype")]
print(f"Top10 类型标签齐备: {'✅' if not weak else '❌ ' + str(weak)}")

print()
print("====== 测试3a: 数据准确性抽查(源头→候选→看板三层一致) ======")
top1 = meta["items"][0]
src = next(r for r in cand if r["id"] == top1["id"])
same = src["engagement_summary"] == top1["engagement_summary"] and \
    src["final_score"] == top1["final_score"]
print(f"Top1: {top1['title'][:45]}")
print(f"  候选={src['engagement_summary']} / 看板={top1['engagement_summary']} / 分数{top1['final_score']}",
      "✅" if same else "❌")

wb = openpyxl.load_workbook(f"{DAY_DIR}/榜单.xlsx")
ws = wb["Top20"]
xlsx_ok = all(ws.cell(r + 2, 4).value == meta["items"][r]["ctype"]
              for r in range(min(20, len(meta["items"]))))
print(f"Excel Top20 行数: {ws.max_row - 1} | 类型列与meta一致: {'✅' if xlsx_ok else '❌'}")

print()
print("====== 测试3b: 看板HTML质量 ======")
html = open("dashboard/index.html", encoding="utf-8").read()
bad = [w for w in ["None", "nan", "undefined", "null"]
       if re.search(r">(\s*)" + w + r"(\s*)<", html)]
cards = len(re.findall(r'<article class="item">', html))
links = html.count('target="_blank"')
whys = len(re.findall(r'class="why"', html))
tabs = len(re.findall('name="day"', html))
print(f"占位残留: {len(bad)} | 卡片: {cards} | 外链: {links} | why: {whys} | 日期tab: {tabs}",
      "✅" if not bad and cards == 20 and whys == 20 else "❌")
print(f"环形图: {'✅' if 'conic-gradient' in html else '❌'} | 分数柱状图: "
      f"{'✅' if 'Top 综合分排行' in html else '❌'}")
print(f"最新日期在页面: {'✅' if DATE in html else '❌'}")

print()
print("====== 测试4: 长期记忆与热度表 ======")
mem_lines = [l for l in open("data/longterm_memory.jsonl", encoding="utf-8") if l.strip()]
print(f"记忆表 {len(mem_lines)} 行（≥当日候选 {len(cand)}）:",
      "✅" if len(mem_lines) >= len(cand) else "❌")
import json as _json  # noqa: E402
heat = _json.loads(open("data/topic_heat.json", encoding="utf-8").read())
bad_t = [t for t, h in heat["topics"].items()
         if not h.get("first_seen") or "status" not in h or "daily" not in h]
print(f"热度表 {len(heat['topics'])} 话题 | 字段缺失: {len(bad_t)}",
      "✅" if not bad_t else bad_t[:3])
hot = meta.get("hot_topics", [])
hot_ok = all({"topic", "today_count", "reason"} <= set(h) for h in hot)
print(f"升温领域标记 {len(hot)} 条，结构完整: {'✅' if hot_ok else '❌'}")
hot_in_dash = hot and hot[0]["topic"] in html
print(f"看板含升温面板: {'✅' if hot_in_dash else ('—今日无升温' if not hot else '❌')}")

print()
print("====== 测试3c: 线上页面与本地一致 ======")
import urllib.request  # noqa: E402
req = urllib.request.Request(
    "https://lj1326644640-blip.github.io/ai-dashboard/",
    headers={"User-Agent": "Mozilla/5.0"})
live = urllib.request.urlopen(req, timeout=20).read().decode("utf-8")
checks = ["AI 爆款看板", "class=\"why\"", "conic-gradient", DATE]
for c in checks:
    print(f"  线上含 '{c}':", "✅" if c in live else "❌")
top1_title_in_live = top1["title"][:30] in live
print(f"  线上=本地数据版本: {'✅' if top1_title_in_live else '❌ (需重新发布)'}")
