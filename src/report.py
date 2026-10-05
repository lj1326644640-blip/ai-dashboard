"""导出：Excel 榜单（打分明细）+ Markdown 日报 + 看板元数据。"""
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from schema import fmt_local

_TZ_HOURS = [8]  # 运行时由 export.set_tz 注入，避免层层传参


def set_tz(offset_hours: int):
    _TZ_HOURS[0] = offset_hours

COLS = [
    ("排名", 6), ("总分", 8), ("平台", 12), ("类型", 8), ("标题", 50), ("链接", 44),
    ("作者", 16), ("发布时间", 12), ("发布至今", 9), ("互动数据", 30),
    ("热度分", 8), ("速度分", 8), ("相关分", 8), ("时效分", 8),
    ("LLM相关", 8), ("LLM价值", 8), ("为什么火", 60),
]


def export_xlsx(ranked: list[dict], all_candidates: list[dict], out: Path, cfg: dict,
                hot: list[dict] | None = None):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Top{cfg['top_n']}"
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for c, (name, width) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width

    for i, r in enumerate(ranked, 2):
        row = [
            i - 1, r["final_score"], r["platform_name"], r.get("ctype", ""), r["title"], r["url"],
            r.get("author", ""), fmt_local(r.get("published"), _TZ_HOURS[0]),
            r["hours_display"], r["engagement_summary"],
            r["s_engagement"], r["s_velocity"], r["s_relevance"], r["s_recency"],
            r.get("llm_relevance"), r.get("llm_value"), r.get("why", ""),
        ]
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=c in (5, 17))
        link_cell = ws.cell(row=i, column=6)
        if r["url"]:
            link_cell.hyperlink = r["url"]
            link_cell.font = Font(color="2563EB", underline="single")
        if r.get("same_story_platforms"):
            note = ws.cell(row=i, column=17)
            note.value = (note.value or "") + f"（同题热传: {', '.join(r['same_story_platforms'])}）"

    last = len(ranked) + 1
    if last >= 2:
        ws.conditional_formatting.add(
            f"B2:B{last}", ColorScaleRule(
                start_type="num", start_value=30, start_color="F87171",
                mid_type="num", mid_value=60, mid_color="FCD34D",
                end_type="num", end_value=85, end_color="34D399"))
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("全部候选")
    for c, (name, width) in enumerate(COLS, 1):
        ws2.cell(row=1, column=c, value=name).font = Font(bold=True)
        ws2.column_dimensions[get_column_letter(c)].width = width
    for i, r in enumerate(all_candidates, 2):
        row = [i - 1, r["final_score"], r["platform_name"], r.get("ctype", ""), r["title"], r["url"],
               r.get("author", ""), fmt_local(r.get("published"), _TZ_HOURS[0]),
               r["hours_display"], r["engagement_summary"], r["s_engagement"],
               r["s_velocity"], r["s_relevance"], r["s_recency"],
               r.get("llm_relevance"), r.get("llm_value"), r.get("why", "")]
        for c, v in enumerate(row, 1):
            ws2.cell(row=i, column=c, value=v)
    ws2.freeze_panes = "A2"

    if hot:
        ws3 = wb.create_sheet("升温领域")
        hot_cols = [("话题", 18), ("今日条数", 9), ("今日均分", 9), ("今日最高分", 10),
                    ("近7天均值条数", 13), ("判定依据", 48), ("代表内容", 50), ("链接", 40)]
        for c, (name, width) in enumerate(hot_cols, 1):
            cell = ws3.cell(row=1, column=c, value=name)
            cell.fill = header_fill
            cell.font = header_font
            ws3.column_dimensions[get_column_letter(c)].width = width
        for i, h in enumerate(hot, 2):
            for c, v in enumerate([h["topic"], h["today_count"], h["today_avg"],
                                   h["today_max"], h["last7_avg_count"], h["reason"],
                                   h["sample_title"], h["sample_url"]], 1):
                cell = ws3.cell(row=i, column=c, value=v)
                cell.alignment = Alignment(vertical="top", wrap_text=c in (6, 7))
            link = ws3.cell(row=i, column=8)
            if h["sample_url"]:
                link.hyperlink = h["sample_url"]
                link.font = Font(color="2563EB", underline="single")
        ws3.freeze_panes = "A2"
    wb.save(out)


def export_md(ranked: list[dict], out: Path, cfg: dict, stats: dict, date_str: str,
              hot: list[dict] | None = None):
    by_platform = {}
    for r in ranked:
        by_platform[r["platform_name"]] = by_platform.get(r["platform_name"], 0) + 1
    plat_line = " · ".join(f"{k}×{v}" for k, v in sorted(by_platform.items()))

    lines = [
        f"# AI 爆款日报 {date_str}",
        "",
        f"> 候选 {stats.get('kept', len(ranked))} 条 → 精选 Top {len(ranked)} | 来源分布：{plat_line or '-'}  ",
        f"> 采集窗口：近{cfg['window_hours']}小时 | 通道：{stats.get('channels', '-')}  ",
        f"> 打分：热度40% + 速度25% + AI相关25% + 价值10%（0-100）",
        "",
    ]
    if hot:
        lines += ["## 🔥 今日升温领域", ""]
        for h in hot:
            lines.append(f"- **{h['topic']}**：今日 {h['today_count']} 条，"
                         f"均分 {h['today_avg']}，最高 {h['today_max']} —— {h['reason']}")
            if h.get("sample_title"):
                lines.append(f"  - 代表：《{h['sample_title'][:60]}》")
        lines.append("")
    medals = ["🥇", "🥈", "🥉"]
    for i, r in enumerate(ranked):
        medal = medals[i] if i < 3 else f"**{i + 1}.**"
        lines += [
            f"### {medal} {r['title']}",
            f"- **平台**：{r['platform_name']} | **类型**：{r.get('ctype', '-')} | **总分 {r['final_score']}** | "
            f"{r['engagement_summary']} | 发布 {fmt_local(r.get('published'), _TZ_HOURS[0])}（{r['hours_display']}前）",
            f"- **链接**：{r['url']}",
            f"- **为什么火**：{r.get('why', '-')}",
        ]
        if r.get("same_story_platforms"):
            lines.append(f"- **同题热传**：{', '.join(r['same_story_platforms'])}（多平台同时发酵）")
        lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")


def write_meta(out: Path, ranked: list[dict], stats: dict, hot: list[dict] | None = None):
    """看板数据源：富字段（含发布时间本地串、分项分、why），够 dashboard 直接渲染。"""
    meta = {
        "top_n": len(ranked),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "stats": stats,
        "hot_topics": hot or [],
        "items": [{
            "rank": i,
            "id": r["id"],
            "platform": r["platform"],
            "platform_name": r["platform_name"],
            "ctype": r.get("ctype", ""),
            "title": r["title"],
            "url": r["url"],
            "author": r.get("author", ""),
            "published_local": fmt_local(r.get("published"), _TZ_HOURS[0]),
            "hours_display": r["hours_display"],
            "engagement_summary": r["engagement_summary"],
            "final_score": r["final_score"],
            "s_engagement": r["s_engagement"],
            "s_velocity": r["s_velocity"],
            "s_relevance": r["s_relevance"],
            "s_recency": r["s_recency"],
            "llm_relevance": r.get("llm_relevance"),
            "llm_value": r.get("llm_value"),
            "why": r.get("why", ""),
            "same_story_platforms": r.get("same_story_platforms", []),
        } for i, r in enumerate(ranked, 1)],
    }
    out.write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
